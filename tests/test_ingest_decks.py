"""Tests for reading a PowerPoint deck into one chunk per slide.

Two kinds of deck are read here. Decks built by hand out of blank slides and
free text boxes, which is what a deck written by a tool rather than from a
template looks like, and the three decks committed to ``data/office``, which
are what the corpus actually holds. The built ones pin the edges — an empty
slide, a title placeholder, a table, an author who is nobody — and the
committed ones prove the reader works on the files it was written for.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from corpus_query.ingest.decks import DECK_SUFFIX, NOTES_LABEL, read_deck
from corpus_query.ingest.pipeline import READERS, ingest_file, reader_for
from corpus_query.ingest.reader import IngestError
from corpus_query.store.kinds import PPTX, SLIDE
from tests.conftest import REPO_ROOT

#: The committed decks, by the slug they are ingested under.
DECKS = (
    "q1-board-review",
    "fieldsense-200-launch-campaign-plan",
    "brannock-refinery-account-recovery-briefing",
)

BLANK_LAYOUT = 6
TITLE_AND_CONTENT_LAYOUT = 1


@pytest.fixture
def deck_paths() -> list[Path]:
    """Return the committed decks, found without depending on the cwd."""
    return [REPO_ROOT / "data" / "office" / f"{slug}{DECK_SUFFIX}" for slug in DECKS]


@pytest.fixture
def build_deck(tmp_path, roster_path):
    """Return a factory that writes a small deck and reads it back.

    Slides are given as ``(title, body lines, notes)``. A title is drawn as a
    free text box in a heading-sized font rather than into a placeholder,
    which is how the committed decks are built.
    """

    def factory(slides, author="Priya", name="deck", **properties):
        presentation = Presentation()
        for title, body, notes in slides:
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[BLANK_LAYOUT]
            )
            if title:
                _text_box(slide, title, top=0.5, size=30)
            if body:
                _text_box(slide, "\n".join(body), top=2.0, size=15)
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
        core = presentation.core_properties
        core.author = author
        core.modified = properties.pop("modified", datetime.datetime(2026, 4, 2))
        for field, value in properties.items():
            setattr(core, field, value)
        path = tmp_path / f"{name}{DECK_SUFFIX}"
        presentation.save(str(path))
        return path

    def read(*args, **kwargs):
        return read_deck(factory(*args, **kwargs), roster_path=roster_path)

    read.path = factory
    return read


def _text_box(slide, text, top, size):
    """Draw one text box on a slide, in a given font size."""
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(9), Inches(1.2))
    frame = box.text_frame
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    return box


def test_the_pipeline_dispatches_pptx_to_the_deck_reader():
    assert READERS[DECK_SUFFIX] is read_deck
    assert reader_for(Path("data/office/anything.PPTX")) is read_deck


def test_a_deck_is_one_chunk_per_slide(build_deck):
    document = build_deck(
        [
            ("Q1 at a Glance", ["Revenue $1.14M"], "We closed the quarter at 91%."),
            ("Pipeline by Stage", ["Commit $367K"], "Coverage is tighter."),
        ]
    )

    assert document.units == 2
    assert [chunk.ordinal for chunk in document.chunks] == [0, 1]
    assert {chunk.kind for chunk in document.chunks} == {SLIDE}
    assert document.source_kind == PPTX


def test_a_slide_carries_its_title_body_and_notes_in_one_chunk(build_deck):
    document = build_deck(
        [("Q1 at a Glance", ["Revenue $1.14M"], "We closed the quarter at 91%.")]
    )

    (chunk,) = document.chunks
    assert "Q1 at a Glance" in chunk.text
    assert "Revenue $1.14M" in chunk.text
    assert NOTES_LABEL in chunk.text
    assert "We closed the quarter at 91%." in chunk.text


def test_a_slides_span_is_its_own_number(build_deck):
    document = build_deck(
        [
            ("Cover", [], ""),
            ("Second", ["A figure"], ""),
        ]
    )

    assert [(chunk.span_start, chunk.span_end) for chunk in document.chunks] == [
        (1, 1),
        (2, 2),
    ]


def test_the_location_names_the_slide_number_and_its_title(build_deck):
    document = build_deck(
        [("Cover", [], ""), ("Root Cause: Thermal Drift", ["±1.8% at 60°C"], "")]
    )

    assert document.chunks[1].location == "slide 2 — Root Cause: Thermal Drift"


def test_a_slide_with_no_text_is_skipped_rather_than_stored_empty(build_deck):
    document = build_deck(
        [
            ("Cover", [], ""),
            (None, [], ""),
            ("Closing", ["Decisions needed"], ""),
        ]
    )

    assert document.units == 3
    assert [chunk.location for chunk in document.chunks] == [
        "slide 1 — Cover",
        "slide 3 — Closing",
    ]
    assert [chunk.ordinal for chunk in document.chunks] == [0, 1]


def test_a_slide_with_only_notes_is_kept(build_deck):
    document = build_deck([(None, [], "The presenter says the whole claim here.")])

    (chunk,) = document.chunks
    assert chunk.location == "slide 1"
    assert "the whole claim" in chunk.text


def test_a_deck_with_no_text_at_all_is_refused(build_deck):
    with pytest.raises(IngestError, match="no text on any slide"):
        build_deck([(None, [], "")])


def test_a_title_placeholder_is_used_when_the_deck_has_one(tmp_path, roster_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[TITLE_AND_CONTENT_LAYOUT]
    )
    slide.shapes.title.text = "Pipeline by Stage"
    slide.placeholders[1].text = "Commit $367K"
    presentation.core_properties.author = "Priya"
    presentation.core_properties.modified = datetime.datetime(2026, 4, 2)
    path = tmp_path / f"placeholders{DECK_SUFFIX}"
    presentation.save(str(path))

    (chunk,) = read_deck(path, roster_path=roster_path).chunks

    assert chunk.location == "slide 1 — Pipeline by Stage"
    assert chunk.text.count("Pipeline by Stage") == 1


def test_a_small_kicker_line_is_not_mistaken_for_the_title(tmp_path, roster_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    _text_box(slide, "BOARD OF DIRECTORS", top=2.0, size=14)
    _text_box(slide, "Q1 2026 Board Review", top=2.5, size=44)
    presentation.core_properties.author = "Priya"
    presentation.core_properties.modified = datetime.datetime(2026, 4, 2)
    path = tmp_path / f"cover{DECK_SUFFIX}"
    presentation.save(str(path))

    (chunk,) = read_deck(path, roster_path=roster_path).chunks

    assert chunk.location == "slide 1 — Q1 2026 Board Review"
    assert "BOARD OF DIRECTORS" in chunk.text


def test_a_table_is_read_one_row_per_line(tmp_path, roster_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK_LAYOUT])
    _text_box(slide, "Pipeline by Stage", top=0.5, size=30)
    table = slide.shapes.add_table(
        2, 2, Inches(0.6), Inches(2), Inches(6), Inches(1.5)
    ).table
    table.cell(0, 0).text = "Stage"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Commit"
    table.cell(1, 1).text = "$367K"
    presentation.core_properties.author = "Priya"
    presentation.core_properties.modified = datetime.datetime(2026, 4, 2)
    path = tmp_path / f"table{DECK_SUFFIX}"
    presentation.save(str(path))

    (chunk,) = read_deck(path, roster_path=roster_path).chunks

    assert "Stage | Value" in chunk.text
    assert "Commit | $367K" in chunk.text


def test_the_document_is_dated_by_its_core_properties(build_deck):
    document = build_deck(
        [("Cover", ["A figure"], "")], modified=datetime.datetime(2026, 2, 10, 9, 30)
    )

    assert document.document_date == "2026-02-10"


def test_the_title_falls_back_to_the_first_slide_title(build_deck):
    document = build_deck([("Q1 at a Glance", ["A figure"], "")], title="")

    assert document.title == "Q1 at a Glance"


def test_a_core_properties_title_wins(build_deck):
    document = build_deck(
        [("Q1 at a Glance", ["A figure"], "")], title="Q1 2026 Board Review"
    )

    assert document.title == "Q1 2026 Board Review"


def test_the_author_resolves_to_one_person_on_the_roster(build_deck):
    document = build_deck([("Cover", ["A figure"], "")], author="priya")

    assert document.author == "Priya"


def test_an_author_who_is_not_on_the_roster_is_refused(build_deck):
    with pytest.raises(IngestError, match="not on the roster"):
        build_deck([("Cover", ["A figure"], "")], author="Gwendolyn")


def test_a_deck_with_no_author_is_refused(build_deck):
    with pytest.raises(IngestError, match="names no author"):
        build_deck([("Cover", ["A figure"], "")], author="")


def test_a_file_that_is_not_a_deck_is_refused(tmp_path, roster_path):
    path = tmp_path / f"not-a-deck{DECK_SUFFIX}"
    path.write_text("This is not a zip archive at all.", encoding="utf-8")

    with pytest.raises(IngestError, match="Could not open"):
        read_deck(path, roster_path=roster_path)


def test_reading_a_deck_twice_gives_the_same_chunks(build_deck, roster_path):
    path = build_deck.path([("Cover", ["A figure"], "Said out loud.")])

    first = read_deck(path, roster_path=roster_path)
    second = read_deck(path, roster_path=roster_path)

    assert first.chunks == second.chunks


def test_the_committed_decks_read_with_an_author_a_date_and_slides(
    deck_paths, roster_path
):
    for path in deck_paths:
        document = read_deck(path, roster_path=roster_path)

        assert document.source_kind == PPTX
        assert document.author
        assert document.attendees == ()
        assert document.document_date.startswith("2026-")
        assert 12 <= document.units <= 18, path.name
        assert len(document.chunks) == document.units, path.name


def test_every_committed_slide_has_a_title_in_its_location(deck_paths, roster_path):
    for path in deck_paths:
        for chunk in read_deck(path, roster_path=roster_path).chunks:
            assert chunk.location.startswith(f"slide {chunk.span_start} — "), (
                f"{path.name}: {chunk.location}"
            )


def test_most_committed_slides_carry_speaker_notes(deck_paths, roster_path):
    for path in deck_paths:
        chunks = read_deck(path, roster_path=roster_path).chunks

        with_notes = [chunk for chunk in chunks if NOTES_LABEL in chunk.text]
        assert len(with_notes) > len(chunks) / 2, path.name


def test_a_committed_deck_ingests_into_the_store(store, deck_paths):
    path = next(p for p in deck_paths if p.stem == "q1-board-review")

    result = ingest_file(store, path)

    assert result.source_kind == PPTX
    assert result.unit_name == "slides"
    assert result.chunks == result.units
    row = store.execute(
        "SELECT title, author, document_date FROM documents WHERE slug = ?",
        (result.slug,),
    ).fetchone()
    assert row["author"] == "Priya"
    assert row["title"] == "Q1 2026 Board Review"
    assert row["document_date"] == "2026-04-02"
    (chunks,) = store.execute(
        "SELECT count(*) FROM chunks WHERE document_id = ? AND kind = ?",
        (result.document_id, SLIDE),
    ).fetchone()
    assert chunks == result.chunks
