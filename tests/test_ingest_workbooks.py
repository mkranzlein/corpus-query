"""Tests for reading an Excel workbook into per-sheet windows of rows.

Most of these build a workbook in the test, so the shape under test — a
sheet of a particular length, a percentage column, a formula with nothing
cached behind it — is on the page next to the assertion. The last few read
the workbooks that are committed to ``data/office``, which is where the
claims about the real corpus are checked.
"""

from __future__ import annotations

import datetime
import re
import zipfile

import openpyxl
import pytest

from corpus_query.ingest.chunk import count_words
from corpus_query.ingest.pipeline import ingest_file, reader_for
from corpus_query.ingest.reader import IngestError
from corpus_query.ingest.workbooks import (
    MAX_WINDOW_ROWS,
    MIN_WINDOW_ROWS,
    read_workbook,
    render_cell,
    windows,
)
from corpus_query.store.kinds import ROW_WINDOW, SHEET_SUMMARY, XLSX
from tests.conftest import REPO_ROOT

#: The committed workbooks, and what each one is expected to carry.
OFFICE_DIR = REPO_ROOT / "data" / "office"
COMMITTED = {
    "q1-sales-pipeline.xlsx": "Jamal",
    "gx-7-bom-cost.xlsx": "Renata",
    "support-ticket-sla-log.xlsx": "Theo",
}

HEADER = ["Ticket ID", "Severity", "Opened Date"]


@pytest.fixture
def write_workbook(tmp_path, roster_path):
    """Return a factory that writes a workbook and reads it back.

    The factory takes sheets as ``{name: rows}``, where the first row is the
    header, and returns the parsed document.
    """

    def factory(
        sheets: dict[str, list[list]] | None = None,
        author: str = "Theo",
        name: str = "support-ticket-sla-log",
        modified: datetime.datetime | None = datetime.datetime(2026, 6, 1),
        prepare=None,
    ):
        workbook = openpyxl.Workbook()
        del workbook[workbook.sheetnames[0]]
        for sheet_name, rows in (sheets or {"Resolved Tickets": _rows(12)}).items():
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
        workbook.properties.creator = author
        workbook.properties.modified = modified
        workbook.properties.created = modified
        if prepare is not None:
            prepare(workbook)
        path = tmp_path / f"{name}.xlsx"
        workbook.save(path)
        return read_workbook(path, roster_path=roster_path)

    return factory


def _rows(count: int) -> list[list]:
    """Return a header and ``count`` data rows of tickets."""
    return [
        HEADER,
        *(
            [f"TKT-{4100 + index}", "Low", datetime.datetime(2026, 1, 1)]
            for index in range(count)
        ),
    ]


def _strip_dates(path) -> None:
    """Rewrite a workbook's core properties without either date.

    openpyxl writes the modified date itself on every save and refuses a
    null created date, so a workbook with no dates at all is made by
    editing the saved file rather than by asking openpyxl for one.
    """
    core = "docProps/core.xml"
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        contents = {name: archive.read(name) for name in names}
    text = contents[core].decode("utf-8")
    text = re.sub(r"<dcterms:(created|modified)[^>]*>.*?</dcterms:\1>", "", text)
    contents[core] = text.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, contents[name])


def row_windows(document):
    """Return the document's row windows."""
    return [chunk for chunk in document.chunks if chunk.kind == ROW_WINDOW]


def sheet_summaries(document):
    """Return the document's per-sheet summary chunks."""
    return [chunk for chunk in document.chunks if chunk.kind == SHEET_SUMMARY]


def test_a_workbook_becomes_a_document_with_its_author_and_date(write_workbook):
    document = write_workbook()

    assert document.source_kind == XLSX
    assert document.author == "Theo"
    assert document.attendees == ()
    assert document.document_date == "2026-06-01"
    assert document.units == 12
    assert document.unit_name == "rows"


def test_the_reader_is_registered_for_the_extension(tmp_path):
    assert reader_for(tmp_path / "book.xlsx") is read_workbook


def test_a_window_is_a_markdown_table_with_the_header_repeated(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(30)})

    for chunk in row_windows(document):
        lines = chunk.text.splitlines()
        assert lines[0] == "**Sheet:** Resolved Tickets"
        assert lines[2] == "| Ticket ID | Severity | Opened Date |"
        assert lines[3] == "| --- | --- | --- |"
        assert chunk.word_count == count_words(chunk.text)


def test_every_row_appears_in_exactly_one_window(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(47)})

    cited = [
        line
        for chunk in row_windows(document)
        for line in chunk.text.splitlines()
        if line.startswith("| TKT-")
    ]

    assert len(cited) == 47
    assert len(set(cited)) == 47


def test_a_window_spans_the_rows_its_location_names(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(47)})

    first, second = row_windows(document)[:2]

    # The header is row 1, so the first data row is row 2.
    assert (first.span_start, first.span_end) == (2, 17)
    assert first.location == "Resolved Tickets, rows 2-17"
    assert second.span_start == first.span_end + 1


def test_a_sheet_of_one_row_is_cited_in_the_singular(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(1)})

    assert row_windows(document)[0].location == "Resolved Tickets, row 2"


@pytest.mark.parametrize("count", [21, 40, 41, 59, 60, 61, 200, 201])
def test_windows_of_a_long_sheet_hold_between_ten_and_twenty_rows(count):
    sizes = [end - start for start, end in windows(count)]

    assert sum(sizes) == count
    assert all(MIN_WINDOW_ROWS <= size <= MAX_WINDOW_ROWS for size in sizes)


def test_a_short_sheet_is_one_window(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(4)})

    assert len(row_windows(document)) == 1


def test_a_sheet_with_no_data_rows_has_a_summary_and_no_windows(write_workbook):
    document = write_workbook({"Resolved Tickets": [HEADER]})

    assert row_windows(document) == []
    assert "no data rows" in sheet_summaries(document)[0].text


def test_reading_the_same_workbook_twice_gives_the_same_windows(
    tmp_path, roster_path, write_workbook
):
    first = write_workbook({"Resolved Tickets": _rows(63)})
    second = read_workbook(
        tmp_path / "support-ticket-sla-log.xlsx", roster_path=roster_path
    )

    assert first.chunks == second.chunks


def test_each_sheet_is_chunked_separately(write_workbook):
    document = write_workbook(
        {"Open": _rows(25), "Closed": _rows(12), "Forecast": _rows(3)}
    )

    assert [chunk.location for chunk in sheet_summaries(document)] == [
        "Open",
        "Closed",
        "Forecast",
    ]
    assert [chunk.location for chunk in row_windows(document)] == [
        "Open, rows 2-14",
        "Open, rows 15-26",
        "Closed, rows 2-13",
        "Forecast, rows 2-4",
    ]


def test_a_sheets_summary_comes_before_its_windows(write_workbook):
    document = write_workbook({"Open": _rows(25), "Closed": _rows(12)})

    kinds = [chunk.kind for chunk in document.chunks]

    assert kinds == [SHEET_SUMMARY, ROW_WINDOW, ROW_WINDOW, SHEET_SUMMARY, ROW_WINDOW]
    assert [chunk.ordinal for chunk in document.chunks] == list(range(len(kinds)))


def test_a_sheet_summary_says_what_the_sheet_holds(write_workbook):
    document = write_workbook({"Resolved Tickets": _rows(12)})

    text = sheet_summaries(document)[0].text

    assert "12 data rows (rows 2-13)" in text
    assert "- Severity: one of Low (12)" in text
    assert "- Opened Date: dates from 2026-01-01 to 2026-01-01" in text
    assert "- Ticket ID: text, 12 distinct values" in text


def test_a_sheet_summary_spans_nothing(write_workbook):
    chunk = sheet_summaries(write_workbook())[0]

    assert (chunk.span_start, chunk.span_end) == (None, None)
    assert chunk.location == "Resolved Tickets"


def test_a_number_column_is_summarized_by_its_range(write_workbook):
    document = write_workbook(
        {
            "Costs": [
                ["Line", "Unit Cost (USD)"],
                *([index, index / 4] for index in [4, 1, 9]),
            ]
        }
    )

    assert (
        "- Unit Cost (USD): numbers from 0.25 to 2.25"
        in sheet_summaries(document)[0].text
    )


def test_a_percentage_cell_reads_as_a_percentage(write_workbook):
    def as_percentage(workbook):
        for row in workbook["Deals"].iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = "0%"

    document = write_workbook(
        {"Deals": [["Deal", "Win Probability"], ["OPP-1", 0.75], ["OPP-2", 0.1]]},
        prepare=as_percentage,
    )

    assert "| OPP-1 | 75% |" in row_windows(document)[0].text
    assert "numbers from 10% to 75%" in sheet_summaries(document)[0].text


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, ""),
        (True, "TRUE"),
        (1450, "1450"),
        (1450.0, "1450"),
        (0.657, "0.657"),
        (datetime.datetime(2026, 1, 5), "2026-01-05"),
        (datetime.datetime(2026, 1, 5, 9, 30), "2026-01-05 09:30:00"),
        ("a | b", "a \\| b"),
        ("two\nlines", "two lines"),
    ],
)
def test_cells_are_rendered_for_a_table(value, expected):
    assert render_cell(value) == expected


def test_a_blank_row_is_left_out_and_the_rows_after_it_keep_their_numbers(
    write_workbook,
):
    rows = _rows(4)
    rows.insert(3, [None, None, None])

    document = write_workbook({"Resolved Tickets": rows})

    assert document.units == 4
    assert row_windows(document)[0].location == "Resolved Tickets, rows 2-6"


def test_a_formula_with_no_cached_value_is_refused(write_workbook):
    rows = _rows(3)
    rows.append(["Total", "=COUNTA(A2:A4)", None])

    with pytest.raises(IngestError, match="formula with no cached value"):
        write_workbook({"Resolved Tickets": rows})


def test_a_workbook_by_someone_not_on_the_roster_is_refused(write_workbook):
    with pytest.raises(IngestError, match="not on the roster"):
        write_workbook(author="Mallory")


def test_a_workbook_with_no_author_is_refused(write_workbook):
    with pytest.raises(IngestError, match="no author"):
        write_workbook(author="")


def test_a_workbook_with_no_date_is_refused(tmp_path, roster_path, write_workbook):
    write_workbook()
    path = tmp_path / "support-ticket-sla-log.xlsx"
    _strip_dates(path)

    with pytest.raises(IngestError, match="no date"):
        read_workbook(path, roster_path=roster_path)


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path, roster_path):
    path = tmp_path / "not-a-workbook.xlsx"
    path.write_text("Not a workbook at all.", encoding="utf-8")

    with pytest.raises(IngestError, match="Could not read"):
        read_workbook(path, roster_path=roster_path)


def test_a_title_is_made_from_the_file_name_as_the_workbook_spells_it(write_workbook):
    document = write_workbook(
        {"Resolved Tickets": [["Ticket ID", "SLA Met"], ["TKT-1", "Yes"]]},
        name="support-ticket-sla-log",
    )

    assert document.title == "Support Ticket SLA Log"


def test_the_core_properties_title_wins_when_there_is_one(write_workbook):
    def set_title(workbook):
        workbook.properties.title = "Support tickets, first half"

    document = write_workbook(prepare=set_title)

    assert document.title == "Support tickets, first half"


@pytest.mark.parametrize(("name", "author"), sorted(COMMITTED.items()))
def test_a_committed_workbook_reads_with_its_author_and_windows(name, author):
    document = read_workbook(OFFICE_DIR / name)

    assert document.author == author
    assert document.source_kind == XLSX
    assert len(row_windows(document)) >= 1
    assert all(
        chunk.span_end - chunk.span_start + 1 <= MAX_WINDOW_ROWS
        for chunk in row_windows(document)
    )


def test_the_committed_pipeline_workbook_has_more_than_one_sheet():
    document = read_workbook(OFFICE_DIR / "q1-sales-pipeline.xlsx")

    assert len(sheet_summaries(document)) == 3


def test_a_committed_workbook_ingests_into_the_store(store):
    result = ingest_file(store, OFFICE_DIR / "support-ticket-sla-log.xlsx")

    row = store.execute(
        "SELECT * FROM documents WHERE id = ?", (result.document_id,)
    ).fetchone()
    kinds = {
        row["kind"]
        for row in store.execute(
            "SELECT DISTINCT kind FROM chunks WHERE document_id = ?",
            (result.document_id,),
        )
    }

    assert row["slug"] == "support-ticket-sla-log"
    assert row["source_kind"] == XLSX
    assert row["author"] == "Theo"
    assert kinds == {ROW_WINDOW, SHEET_SUMMARY}
    assert result.unit_name == "rows"


def test_an_ingested_workbook_is_searchable_by_a_value_in_a_row(store):
    from corpus_query.retrieval.lexical import search_lexical

    ingest_file(store, OFFICE_DIR / "support-ticket-sla-log.xlsx")

    hits = search_lexical(store, "TKT-4101")

    assert hits
