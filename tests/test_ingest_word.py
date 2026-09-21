"""Tests for reading Word documents into heading sections.

Nothing here is mocked. The documents are built with python-docx into a
temporary directory, read back with the reader, and checked against what a
citation would claim about them. The three committed documents are read too,
because a chunker that behaves on a fixture and not on the corpus is not
worth much.
"""

from __future__ import annotations

import datetime

import docx
import pytest

from corpus_query.ingest.chunk import count_words
from corpus_query.ingest.pipeline import READERS, document_paths, ingest_file
from corpus_query.ingest.reader import IngestError
from corpus_query.ingest.word import (
    DOCX_SUFFIX,
    parse_word_document,
    read_word_document,
)
from corpus_query.store.kinds import DOCX, DOCX_SECTION
from tests.conftest import REPO_ROOT

#: Where the committed Word documents live.
OFFICE_DIR = REPO_ROOT / "data" / "office"

WHEN = datetime.datetime(2026, 3, 12, 9, 0, tzinfo=datetime.UTC)


@pytest.fixture
def word_file(tmp_path, roster_path):
    """Return a factory that writes a Word document and reads it back."""

    def factory(
        body,
        name: str = "transition-plan.docx",
        author: str = "Devon",
        title: str = "Transition Plan",
        created=WHEN,
        modified=WHEN,
        target_words: int = 250,
        read: bool = True,
    ):
        document = docx.Document()
        for style, text in body:
            if style == "table":
                table = document.add_table(rows=len(text), cols=len(text[0]))
                for row, cells in zip(table.rows, text, strict=True):
                    for cell, value in zip(row.cells, cells, strict=True):
                        cell.text = value
            else:
                document.add_paragraph(text, style=style)
        properties = document.core_properties
        properties.author = author
        properties.last_modified_by = author
        properties.title = title
        _set_date(properties, "created", created)
        _set_date(properties, "modified", modified)
        path = tmp_path / name
        document.save(path)
        if not read:
            return path
        return read_word_document(
            path, target_words=target_words, roster_path=roster_path
        )

    return factory


def _set_date(properties, name: str, value) -> None:
    """Set a core property date, or take it out of the file entirely.

    python-docx will not assign ``None`` to a date property, so a document
    that carries no date at all is made by removing the element — which is
    also what a file written by something other than Word tends to look
    like.
    """
    if value is not None:
        setattr(properties, name, value)
        return
    element = properties._element
    for child in list(element):
        if child.tag.endswith(f"}}{name}"):
            element.remove(child)


def paragraphs(path) -> set[str]:
    """Return every paragraph of a Word document, as text."""
    return {
        paragraph.text.strip()
        for paragraph in docx.Document(str(path)).paragraphs
        if paragraph.text.strip()
    }


def test_the_header_fields_come_from_the_core_properties(word_file):
    document = word_file([("Heading 1", "Scope"), (None, "What this covers.")])

    assert document.source_kind == DOCX
    assert document.title == "Transition Plan"
    assert document.document_date == "2026-03-12"
    assert document.author == "Devon"
    assert document.attendees == ()
    assert document.unit_name == "sections"


def test_a_section_becomes_a_chunk_labelled_with_its_heading_path(word_file):
    document = word_file(
        [
            ("Heading 1", "Transition Risks"),
            (None, "Two things could go wrong."),
            ("Heading 2", "Supplier Qualification Risk"),
            (None, "The second source has not built this board."),
        ]
    )

    first, second = document.chunks
    assert document.units == 2
    assert first.location == "Transition Risks"
    assert second.location == "Transition Risks > Supplier Qualification Risk"
    assert second.kind == DOCX_SECTION
    assert "The second source has not built this board." in second.text


def test_a_deeper_heading_extends_the_path_and_a_shallower_one_replaces_it(word_file):
    document = word_file(
        [
            ("Heading 1", "Thermal Chamber Results"),
            ("Heading 2", "Chamber B, High Ambient"),
            ("Heading 3", "Chamber B, 60 C Soak"),
            (None, "Mean error was 1.62%."),
            ("Heading 1", "Findings"),
            (None, "The workaround holds to 60 C."),
        ]
    )

    deep, shallow = document.chunks
    assert deep.location == (
        "Thermal Chamber Results > Chamber B, High Ambient > Chamber B, 60 C Soak"
    )
    assert shallow.location == "Findings"


def test_a_heading_with_no_body_is_carried_into_the_section_below_it(word_file):
    document = word_file(
        [
            ("Heading 1", "Thermal Chamber Results"),
            ("Heading 2", "Chamber A, Low and Mid Range"),
            (None, "Error stayed inside specification."),
        ]
    )

    [chunk] = document.chunks
    assert chunk.text.splitlines() == [
        "Thermal Chamber Results",
        "Chamber A, Low and Mid Range",
        "Error stayed inside specification.",
    ]


def test_text_written_above_the_first_heading_is_filed_under_the_title(word_file):
    document = word_file(
        [
            (None, "This plan covers board assembly only."),
            ("Heading 1", "Current State"),
            (None, "One contract manufacturer builds every board."),
        ]
    )

    preamble, _ = document.chunks
    assert preamble.location == "Transition Plan"
    assert preamble.text == "This plan covers board assembly only."


def test_the_chunk_spans_say_where_in_the_file_the_text_sits(word_file):
    document = word_file(
        [
            ("Heading 1", "Scope"),
            (None, "What this covers."),
            (None, "What it does not."),
            ("Heading 1", "Method"),
            (None, "Three chambers, four units."),
        ]
    )

    scope, method = document.chunks
    assert (scope.span_start, scope.span_end) == (0, 2)
    assert (method.span_start, method.span_end) == (3, 4)


def test_a_table_is_kept_where_it_was_written_and_read_as_rows(word_file):
    document = word_file(
        [
            ("Heading 1", "Preliminary Quotes"),
            (None, "Three assessment bodies quoted."),
            ("table", [["Body", "Quote"], ["Keystone", "48,000"]]),
            (None, "All three quotes exclude travel."),
        ]
    )

    [chunk] = document.chunks
    assert chunk.text.splitlines() == [
        "Preliminary Quotes",
        "Three assessment bodies quoted.",
        "| Body | Quote |",
        "| Keystone | 48,000 |",
        "All three quotes exclude travel.",
    ]


def test_a_long_section_splits_on_a_paragraph_boundary_and_keeps_its_path(word_file):
    paragraph = " ".join(["word"] * 40)
    document = word_file(
        [
            ("Heading 1", "Findings"),
            ("Heading 2", "Unit-to-Unit Variation"),
            *[(None, f"{index} {paragraph}") for index in range(3)],
        ],
        target_words=100,
    )

    first, second = document.chunks
    assert document.units == 1, "a split is one section, not two"
    assert first.location == second.location == "Findings > Unit-to-Unit Variation"
    assert first.word_count <= 100
    assert first.span_end + 1 == second.span_start
    assert "Unit-to-Unit Variation" not in second.text, (
        "the heading belongs to the first part only, not stored twice"
    )


def test_a_paragraph_longer_than_a_chunk_is_not_cut_in_half(word_file):
    paragraph = " ".join(["word"] * 300)
    document = word_file(
        [("Heading 1", "Yield Dip During Ramp"), (None, paragraph)], target_words=100
    )

    [chunk] = document.chunks
    assert chunk.text.endswith(paragraph)


def test_the_word_count_is_the_words_of_the_chunk(word_file):
    document = word_file([("Heading 1", "Scope"), (None, "Four words go here.")])

    [chunk] = document.chunks
    assert chunk.word_count == count_words(chunk.text) == 5


def test_an_author_who_is_not_on_the_roster_is_refused(word_file):
    with pytest.raises(IngestError, match="not on the roster"):
        word_file([("Heading 1", "Scope"), (None, "Body.")], author="Gwen")


def test_a_document_with_no_author_is_refused(word_file):
    with pytest.raises(IngestError, match="names no author"):
        word_file([("Heading 1", "Scope"), (None, "Body.")], author="")


def test_the_author_is_spelled_the_way_the_roster_spells_it(word_file):
    document = word_file([("Heading 1", "Scope"), (None, "Body.")], author="devon")

    assert document.author == "Devon"


def test_a_document_with_no_dates_at_all_is_refused(word_file):
    with pytest.raises(IngestError, match="no created or modified date"):
        word_file(
            [("Heading 1", "Scope"), (None, "Body.")], created=None, modified=None
        )


def test_a_document_that_was_never_modified_is_dated_by_its_creation(word_file):
    document = word_file(
        [("Heading 1", "Scope"), (None, "Body.")], created=WHEN, modified=None
    )

    assert document.document_date == "2026-03-12"


def test_a_document_is_dated_by_when_it_was_last_written(word_file):
    document = word_file(
        [("Heading 1", "Scope"), (None, "Body.")],
        created=WHEN,
        modified=datetime.datetime(2026, 4, 15, 9, 0, tzinfo=datetime.UTC),
    )

    assert document.document_date == "2026-04-15"


def test_a_document_with_no_title_property_falls_back_to_its_first_heading(word_file):
    document = word_file(
        [("Heading 1", "Scope of This Plan"), (None, "Body.")], title=""
    )

    assert document.title == "Scope of This Plan"


def test_a_file_that_is_not_a_word_document_is_refused(tmp_path, roster_path):
    path = tmp_path / "notes.docx"
    path.write_text("Just some notes.\n", encoding="utf-8")

    with pytest.raises(IngestError, match="Could not read"):
        read_word_document(path, roster_path=roster_path)


def test_the_pipeline_dispatches_docx_to_this_reader(word_file, store):
    path = word_file([("Heading 1", "Scope"), (None, "What this covers.")], read=False)

    result = ingest_file(store, path)

    assert DOCX_SUFFIX in READERS
    assert result.source_kind == DOCX
    assert result.chunks == 1
    row = store.execute("SELECT author, title FROM documents").fetchone()
    assert (row["author"], row["title"]) == ("Devon", "Transition Plan")


def test_the_default_directories_cover_the_transcripts_and_the_office_files():
    names = [path.name for path in document_paths()]

    assert "xt-9-rev-b-thermal-qualification-report.docx" in names
    assert any(name.endswith(".md") for name in names)


@pytest.mark.parametrize(
    "name",
    [
        "contract-manufacturer-transition-plan.docx",
        "iec-62443-certification-readiness-assessment.docx",
        "xt-9-rev-b-thermal-qualification-report.docx",
    ],
)
def test_a_committed_document_chunks_into_its_sections(name, roster_path):
    path = OFFICE_DIR / name

    document = read_word_document(path, roster_path=roster_path)

    assert len(document.chunks) > 5, "a document is not one chunk"
    assert all(chunk.kind == DOCX_SECTION for chunk in document.chunks)
    assert any(" > " in chunk.location for chunk in document.chunks)
    assert document.author in {"Sofia", "Callum", "Devon"}


@pytest.mark.parametrize(
    "name",
    [
        "contract-manufacturer-transition-plan.docx",
        "iec-62443-certification-readiness-assessment.docx",
        "xt-9-rev-b-thermal-qualification-report.docx",
    ],
)
def test_every_line_of_a_chunk_is_findable_in_the_document(name, roster_path):
    path = OFFICE_DIR / name
    written = paragraphs(path)

    document = read_word_document(path, roster_path=roster_path)

    quoted = {
        line
        for chunk in document.chunks
        for line in chunk.text.splitlines()
        if not line.startswith("|")
    }
    assert quoted <= written


def test_the_committed_documents_use_three_heading_levels(roster_path):
    parsed = parse_word_document(
        OFFICE_DIR / "xt-9-rev-b-thermal-qualification-report.docx",
        roster_path=roster_path,
    )

    levels = {block.heading_level for block in parsed.blocks}
    assert {1, 2, 3} <= levels
