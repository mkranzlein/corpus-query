"""Tests for converting legacy office files before reading them.

No test here invokes LibreOffice. There are no legacy fixtures to convert
and committing one is out of scope, so the converter is stubbed throughout:
these tests check the dispatch, the delegation to the modern reader, that
the legacy file's own name survives into the document, and that a missing
binary or a failed conversion each raise a clear ``IngestError``. The one
check that needs LibreOffice for real is done by hand and written up in the
pull request instead.
"""

from __future__ import annotations

import datetime

import docx
import pytest

from corpus_query.ingest import legacy
from corpus_query.ingest.pipeline import READERS, ingest_file
from corpus_query.ingest.reader import IngestError
from corpus_query.ingest.word import DOCX_SUFFIX, read_word_document
from corpus_query.store.db import connect

WHEN = datetime.datetime(2026, 3, 12, 9, 0, tzinfo=datetime.UTC)


@pytest.fixture
def store():
    """Return an open, empty document store."""
    connection = connect(":memory:")
    yield connection
    connection.close()


@pytest.fixture
def modern_docx(tmp_path, roster_path):
    """Write a real ``.docx`` file that stands in for a converted document."""

    def factory(name: str = "transition-plan.docx"):
        document = docx.Document()
        document.add_paragraph("Background", style="Heading 1")
        document.add_paragraph("The plan is on schedule.")
        properties = document.core_properties
        properties.author = "Devon"
        properties.last_modified_by = "Devon"
        properties.created = WHEN
        properties.modified = WHEN
        path = tmp_path / "converted" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)
        return path

    return factory


@pytest.fixture
def legacy_file(tmp_path):
    """Return a factory that writes a placeholder legacy file.

    Its bytes are never read by anything in these tests: conversion is
    always stubbed, so a real legacy binary is not needed to exercise the
    dispatch, the delegation, or the naming.
    """

    def factory(name: str = "transition-plan.doc"):
        path = tmp_path / name
        path.write_bytes(b"not a real .doc file")
        return path

    return factory


def test_doc_ppt_and_xls_are_registered() -> None:
    assert set(READERS) >= {".doc", ".ppt", ".xls"}


def test_a_doc_file_delegates_to_the_word_reader(
    monkeypatch, store, legacy_file, modern_docx
):
    source = legacy_file()
    converted = modern_docx()
    monkeypatch.setattr(
        legacy, "convert", lambda path, target_suffix, outdir: converted
    )

    result = ingest_file(store, source)

    row = store.execute("SELECT * FROM documents").fetchone()
    assert result.source_kind == row["source_kind"]
    assert row["title"] == "Background"


def test_the_legacy_name_survives_into_the_slug_and_source_path(
    monkeypatch, store, legacy_file, modern_docx
):
    source = legacy_file("q3-report.doc")
    converted = modern_docx("q3-report.docx")
    monkeypatch.setattr(
        legacy, "convert", lambda path, target_suffix, outdir: converted
    )

    result = ingest_file(store, source)

    assert result.slug == "q3-report"
    row = store.execute("SELECT slug, source_path FROM documents").fetchone()
    assert row["slug"] == "q3-report"
    assert row["source_path"] == str(source)


def test_the_legacy_name_survives_when_the_modern_title_falls_back_to_it(
    monkeypatch, store, legacy_file, modern_docx
):
    # The converted file's stem matches the legacy file's own stem, because
    # that is what LibreOffice writes, so a title that falls back to the
    # filename still names the legacy document rather than a temporary path.
    source = legacy_file("q3-report.doc")
    converted = modern_docx("q3-report.docx")
    document = docx.Document(str(converted))
    document.core_properties.title = ""
    for paragraph in document.paragraphs:
        paragraph.style = document.styles["Normal"]
    document.save(converted)
    monkeypatch.setattr(
        legacy, "convert", lambda path, target_suffix, outdir: converted
    )

    ingest_file(store, source)

    row = store.execute("SELECT title FROM documents").fetchone()
    assert row["title"] == "q3-report"


def test_converting_reader_calls_the_modern_reader_directly(
    monkeypatch, legacy_file, modern_docx
):
    source = legacy_file()
    converted = modern_docx()
    calls = []

    def fake_modern_reader(path, target_words):
        calls.append((path, target_words))
        return read_word_document(path, target_words=target_words)

    monkeypatch.setattr(
        legacy, "convert", lambda path, target_suffix, outdir: converted
    )
    reader = legacy.converting_reader(DOCX_SUFFIX, fake_modern_reader)
    document = reader(source, 250)

    assert calls == [(converted, 250)]
    assert document.title == "Background"


def test_a_missing_binary_is_a_clear_ingest_error(monkeypatch, legacy_file, tmp_path):
    source = legacy_file()
    monkeypatch.setattr(legacy, "find_soffice", lambda: None)

    with pytest.raises(IngestError) as excinfo:
        legacy.convert(source, DOCX_SUFFIX, tmp_path)

    message = str(excinfo.value)
    assert str(source) in message
    assert "LibreOffice" in message
    assert "install" in message.lower()


def test_ingesting_without_libreoffice_reports_and_skips_one_file(
    monkeypatch, store, legacy_file
):
    monkeypatch.setattr(legacy, "find_soffice", lambda: None)
    source = legacy_file()

    with pytest.raises(IngestError) as excinfo:
        ingest_file(store, source)

    assert str(source) in str(excinfo.value)
    assert not store.execute("SELECT id FROM documents").fetchall()


def test_a_failed_conversion_is_an_ingest_error(monkeypatch, legacy_file, tmp_path):
    source = legacy_file()

    class FakeResult:
        returncode = 1
        stdout = ""
        stderr = "Error: source file could not be loaded"

    monkeypatch.setattr(legacy, "find_soffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(legacy.subprocess, "run", lambda *a, **k: FakeResult())

    with pytest.raises(IngestError) as excinfo:
        legacy.convert(source, DOCX_SUFFIX, tmp_path)

    message = str(excinfo.value)
    assert str(source) in message
    assert "could not be loaded" in message


def test_a_conversion_with_no_output_is_an_ingest_error(
    monkeypatch, legacy_file, tmp_path
):
    source = legacy_file()

    class FakeResult:
        returncode = 0
        stdout = ""
        stderr = ""

    monkeypatch.setattr(legacy, "find_soffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(legacy.subprocess, "run", lambda *a, **k: FakeResult())

    with pytest.raises(IngestError) as excinfo:
        legacy.convert(source, DOCX_SUFFIX, tmp_path)

    assert str(source) in str(excinfo.value)


def test_find_soffice_prefers_path(monkeypatch, tmp_path):
    on_path = tmp_path / "soffice"
    monkeypatch.setattr(legacy.shutil, "which", lambda name: str(on_path))
    assert legacy.find_soffice() == on_path


def test_find_soffice_falls_back_to_the_macos_app(monkeypatch, tmp_path):
    monkeypatch.setattr(legacy.shutil, "which", lambda name: None)
    fake_app = tmp_path / "soffice"
    fake_app.write_text("", encoding="utf-8")
    monkeypatch.setattr(legacy, "MACOS_SOFFICE_PATH", fake_app)
    assert legacy.find_soffice() == fake_app


def test_find_soffice_is_none_when_neither_exists(monkeypatch, tmp_path):
    monkeypatch.setattr(legacy.shutil, "which", lambda name: None)
    monkeypatch.setattr(legacy, "MACOS_SOFFICE_PATH", tmp_path / "absent")
    assert legacy.find_soffice() is None
