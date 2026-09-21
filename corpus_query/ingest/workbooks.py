"""Reading an Excel workbook and cutting each sheet into windows of rows.

This is one format reader among several. It reads a ``.xlsx`` file with
openpyxl and cuts it along the seams a spreadsheet has: sheets, and rows
within a sheet.

Row 1 of every sheet is its header. Each window is a run of 10 to 20 data
rows rendered as a markdown table with that header repeated on top, so a
window read on its own still says what each column is. A window's span is
the rows' own spreadsheet row numbers — the header is row 1, the first data
row is row 2 — so the ``location`` a citation shows, ``Open Opportunities,
rows 2-16``, is what someone opening the file sees in the margin.

Windows do not overlap. A turn in a meeting leans on the one before it, which
is why transcript windows share a turn at each boundary; a row in a sheet
stands on its own, and repeating one would only count it twice.

Each sheet also gets one summary chunk, ahead of its windows, that says what
the sheet holds: how many rows, which columns, and what is in each column.
It is derived from the cells, not written by a model, so it is as
deterministic as the windows and costs nothing to produce. It is there for
the questions a window cannot answer on its own — how many tickets there
are, what date range a sheet covers, who owns the rows — and is recorded as
``sheet_summary``, a kind that carries no span.

Cell values are read with ``data_only=True``, which returns the value a
formula last computed rather than the formula's text. That value exists only
if whatever saved the file cached it; a workbook written by a library and
never opened in a spreadsheet application has formulas with nothing behind
them. Such a cell is refused rather than rendered as an empty one, since a
blank where a total should be reads as a total of nothing.

A workbook has one author, taken from its core properties, and that author
has to be someone on the roster. A name nobody on the roster answers to is
refused rather than stored, because a document attributed to no one in
particular cannot answer a question about who wrote what.

Nothing here is random or time-dependent, so reading the same workbook twice
gives the same chunks.
"""

from __future__ import annotations

import datetime
import math
import zipfile
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from corpus_query.ingest.chunk import TARGET_WORDS, Chunk, count_words
from corpus_query.ingest.reader import IngestError, ReadDocument, span_location
from corpus_query.store.kinds import ROW_WINDOW, SHEET_SUMMARY, XLSX
from corpus_query.transcripts.roster import (
    DEFAULT_ROSTER_FILE,
    RosterError,
    first_names,
    read_roster,
)

#: The extension this reader claims.
XLSX_SUFFIX = ".xlsx"

#: The fewest and most data rows a window holds. A sheet is split into as few
#: windows as keep each one at or under the maximum, and the rows are shared
#: out as evenly as they go, so every window of a sheet longer than the
#: maximum holds at least the minimum. A sheet shorter than the minimum is
#: one window of whatever it has.
MIN_WINDOW_ROWS = 10
MAX_WINDOW_ROWS = 20

#: What a row window's citation label calls the units it spans.
ROW_UNIT = "row"

#: The spreadsheet row the header sits on. Data starts on the next one.
HEADER_ROW = 1

#: A text column with at most this many distinct values is described by
#: listing them, as a category; one with more is described as free text.
MAX_LISTED_VALUES = 10

#: The roster, found from the package rather than the working directory, so
#: ingesting from anywhere resolves authors against the same file.
ROSTER_FILE = Path(__file__).resolve().parents[2] / DEFAULT_ROSTER_FILE

#: Where a workbook's core properties live inside the file.
CORE_PROPERTIES = "docProps/core.xml"

#: A row of cell values, as openpyxl hands them back.
type Row = tuple[object, ...]


@dataclass(frozen=True)
class Sheet:
    """One sheet of a workbook, reduced to its header and its data rows."""

    name: str
    header: tuple[str, ...]
    """Column labels, rendered."""

    rows: tuple[tuple[int, tuple[str, ...]], ...]
    """Each data row's spreadsheet row number and its rendered cells. Rows
    with nothing in them are left out, but the rows around them keep their
    own numbers."""

    values: tuple[Row, ...]
    """The same data rows' raw values, in the same order, for describing
    the columns."""

    formats: tuple[str, ...]
    """Each column's number format, taken from its first data row, so a
    column is described in the units it is displayed in."""


def read_workbook(
    path: Path,
    target_words: int = TARGET_WORDS,
    roster_path: Path = ROSTER_FILE,
) -> ReadDocument:
    """Read one workbook into a document ready to be written.

    Args:
        path: The workbook to read.
        target_words: Accepted so this matches every other reader, and
            ignored: a window is sized in rows, not words.
        roster_path: The roster the author has to be on.

    Returns:
        The document, with one summary chunk and a run of row windows per
        sheet.

    Raises:
        IngestError: If the file cannot be opened as a workbook, has no
            author or no date in its core properties, names an author who is
            not on the roster, or holds a formula with no cached value.
    """
    del target_words
    values = _load(path, data_only=True)
    formulas = _load(path, data_only=False)

    author = _author(path, values, roster_path)
    document_date = _document_date(path, values)
    sheets = [
        _read_sheet(path, sheet, formulas[sheet.title]) for sheet in values.worksheets
    ]
    title = _title(path, values, sheets)

    chunks: list[Chunk] = []
    for sheet in sheets:
        if sheet is None:
            continue
        chunks.append(_summary_chunk(len(chunks), title, sheet))
        for start, end in windows(len(sheet.rows)):
            chunks.append(_window_chunk(len(chunks), sheet, sheet.rows[start:end]))

    return ReadDocument(
        source_kind=XLSX,
        title=title,
        document_date=document_date,
        author=author,
        attendees=(),
        chunks=tuple(chunks),
        units=sum(len(sheet.rows) for sheet in sheets if sheet is not None),
        unit_name="rows",
    )


def windows(count: int) -> list[tuple[int, int]]:
    """Split a run of rows into windows of 10 to 20.

    Args:
        count: How many data rows the sheet has.

    Returns:
        One half-open ``(start, end)`` pair of row positions per window, in
        order, covering every row exactly once. Sizes differ by at most one,
        larger windows first. Empty when there are no rows.
    """
    if count <= 0:
        return []
    number = math.ceil(count / MAX_WINDOW_ROWS)
    size, extra = divmod(count, number)
    bounds = []
    start = 0
    for index in range(number):
        end = start + size + (1 if index < extra else 0)
        bounds.append((start, end))
        start = end
    return bounds


def render_cell(value: object, number_format: str = "General") -> str:
    """Render one cell's value as table text.

    Dates are ISO dates, with the time only when there is one. A number
    formatted as a percentage is shown as one, since the stored ``0.75`` of
    a cell that reads ``75%`` would be misread as a fraction of a unit.
    Other numbers are shown in full, without the thousands separators or
    currency symbol their format adds: the header says what the unit is,
    and the bare figure is what a search for it will type.

    Args:
        value: The cell's value, as openpyxl read it.
        number_format: The cell's number format.

    Returns:
        The text, safe to put between the pipes of a markdown table.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if "%" in number_format:
            return f"{_number(value * 100)}%"
        return _number(value)
    text = " ".join(str(value).split())
    return text.replace("|", "\\|")


def _number(value: float) -> str:
    """Render a number in full, dropping a ``.0`` that says nothing."""
    if isinstance(value, float):
        value = round(value, 10)
        if value.is_integer():
            return str(int(value))
    return repr(value)


def _load(path: Path, data_only: bool) -> Workbook:
    """Open a workbook, turning any failure into an :class:`IngestError`."""
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise IngestError(f"Could not read {path} as a workbook: {exc}") from exc


def _author(path: Path, workbook: Workbook, roster_path: Path) -> str:
    """Return the workbook's author, checked against the roster.

    Args:
        path: The workbook, for error messages.
        workbook: The opened workbook.
        roster_path: The roster the author has to be on.

    Returns:
        The author's first name, exactly as the roster spells it.

    Raises:
        IngestError: If the author is missing, or is not someone on the
            roster.
    """
    author = (workbook.properties.creator or "").strip()
    if not author:
        raise IngestError(f"{path} has no author in its core properties.")
    try:
        names = first_names(read_roster(roster_path))
    except RosterError as exc:
        raise IngestError(f"Could not check the author of {path}: {exc}") from exc
    if author not in names:
        raise IngestError(
            f"{path} names {author!r} as its author, who is not on the roster "
            f"at {roster_path}. The author has to be one of: {', '.join(names)}."
        )
    return author


def _document_date(path: Path, workbook: Workbook) -> str:
    """Return when the workbook is dated: created, else last modified.

    Created is preferred because it is the date the workbook is about,
    while the modified date is whenever the file was last written — a
    library that opens a workbook and saves it again moves that date to
    today without a cell having changed.

    Args:
        path: The workbook, to check its core properties carry a date at
            all.
        workbook: The opened workbook.

    Returns:
        The date, as ``YYYY-MM-DD``.

    Raises:
        IngestError: If the core properties carry neither date. That is
            checked against the file rather than against what openpyxl
            reports, because openpyxl answers for a workbook with no dates
            in it by handing back the time it was opened, which would date
            every such document today without saying so.
    """
    if not _dated(path):
        raise IngestError(f"{path} has no date in its core properties.")
    properties = workbook.properties
    when = properties.created or properties.modified
    if when is None:
        raise IngestError(f"{path} has no date in its core properties.")
    return when.date().isoformat()


def _dated(path: Path) -> bool:
    """Return whether the file's core properties carry a date at all."""
    try:
        with zipfile.ZipFile(path) as archive:
            core = archive.read(CORE_PROPERTIES).decode("utf-8", errors="replace")
    except OSError, KeyError, zipfile.BadZipFile:
        return False
    return "dcterms:created" in core or "dcterms:modified" in core


def _title(path: Path, workbook: Workbook, sheets: Sequence[Sheet | None]) -> str:
    """Return the workbook's title.

    The core properties' title when it has one. Otherwise one made from the
    file name, whose words carry no capitalization at all: ``q1-sales-
    pipeline``. Each word is spelled the way the workbook itself spells it
    wherever it says it — a sheet named ``ProximaSense X4 BOM`` settles
    three of the four words of that file's name — and capitalized when the
    workbook never does.

    Args:
        path: The workbook's path, whose stem the title is made from.
        workbook: The opened workbook.
        sheets: Its sheets, for the spellings they carry.

    Returns:
        The title.
    """
    title = (workbook.properties.title or "").strip()
    if title:
        return title
    spellings = _spellings(sheets)
    return " ".join(
        spellings.get(word.lower(), word.capitalize())
        for word in path.stem.replace("_", "-").split("-")
        if word
    )


def _spellings(sheets: Sequence[Sheet | None]) -> dict[str, str]:
    """Collect how the workbook spells each word it uses.

    Args:
        sheets: The workbook's sheets.

    Returns:
        The most common capitalized spelling of each word, keyed by that
        word in lower case. A word the workbook only ever writes in lower
        case is left out, so the caller capitalizes it itself.
    """
    counts: dict[str, Counter[str]] = {}
    for sheet in sheets:
        if sheet is None:
            continue
        cells = (cell for _, row in sheet.rows for cell in row)
        for text in [sheet.name, *sheet.header, *cells]:
            for word in _words(text):
                if word.islower():
                    continue
                counts.setdefault(word.lower(), Counter())[word] += 1
    return {word: spellings.most_common(1)[0][0] for word, spellings in counts.items()}


def _words(text: str) -> list[str]:
    """Split text into runs of letters and digits."""
    words: list[str] = []
    current: list[str] = []
    for char in text:
        if char.isalnum():
            current.append(char)
        elif current:
            words.append("".join(current))
            current = []
    if current:
        words.append("".join(current))
    return words


def _read_sheet(path: Path, values: Worksheet, formulas: Worksheet) -> Sheet | None:
    """Read one sheet's header and data rows.

    Args:
        path: The workbook, for error messages.
        values: The sheet as read with cached values.
        formulas: The same sheet as read with formulas, to tell a formula
            with no cached value from a cell that is empty.

    Returns:
        The sheet, or ``None`` if it has nothing in it at all.

    Raises:
        IngestError: If a formula cell has no cached value.
    """
    width = 0
    table: list[tuple[int, list[str], Row, tuple[str, ...]]] = []
    for cells, formula_cells in zip(
        values.iter_rows(), formulas.iter_rows(), strict=True
    ):
        raw = tuple(cell.value for cell in cells)
        for cell, formula_cell in zip(cells, formula_cells, strict=True):
            if formula_cell.data_type == "f" and cell.value is None:
                raise IngestError(
                    f"{path}: cell {cell.coordinate} on sheet {values.title!r} "
                    f"holds a formula with no cached value, so there is no "
                    f"computed value to read. Open the workbook in a "
                    f"spreadsheet application and save it, so the value is "
                    f"stored alongside the formula."
                )
        if all(value is None for value in raw):
            continue
        width = max(
            width, max(i for i, value in enumerate(raw) if value is not None) + 1
        )
        formats = tuple(cell.number_format or "General" for cell in cells)
        rendered = [
            render_cell(cell.value, number_format)
            for cell, number_format in zip(cells, formats, strict=True)
        ]
        table.append((cells[0].row, rendered, raw, formats))

    if not table:
        return None
    header_row, header, _, _ = table[0]
    if header_row != HEADER_ROW:
        raise IngestError(
            f"{path}: sheet {values.title!r} starts on row {header_row}. Row "
            f"{HEADER_ROW} has to be the header."
        )

    def pad(cells: Sequence) -> tuple:
        return tuple(cells[:width]) + ("",) * (width - len(cells))

    data = table[1:]
    return Sheet(
        name=values.title,
        header=pad(header),
        rows=tuple((number, pad(rendered)) for number, rendered, _, _ in data),
        values=tuple(
            tuple(raw[:width]) + (None,) * (width - len(raw)) for _, _, raw, _ in data
        ),
        formats=(
            tuple(data[0][3][:width]) + ("General",) * (width - len(data[0][3]))
            if data
            else ("General",) * width
        ),
    )


def _table(header: Sequence[str], rows: Sequence[Sequence[str]]) -> str:
    """Render a header and rows as a markdown table."""
    lines = [
        _table_row(header),
        _table_row(["---"] * len(header)),
        *(_table_row(row) for row in rows),
    ]
    return "\n".join(lines)


def _table_row(cells: Sequence[str]) -> str:
    """Render one markdown table row."""
    return "| " + " | ".join(cells) + " |"


def _window_chunk(
    ordinal: int, sheet: Sheet, rows: Sequence[tuple[int, tuple[str, ...]]]
) -> Chunk:
    """Assemble one row window.

    The sheet's name heads the text as well as the location. The location
    is what a citation shows, but the text is what is embedded and indexed,
    and a window that does not say which sheet it came from reads the same
    as any other sheet with the same columns.

    Args:
        ordinal: The chunk's position in the document.
        sheet: The sheet the rows belong to.
        rows: The window's rows, with their spreadsheet row numbers.

    Returns:
        The chunk.
    """
    start, end = rows[0][0], rows[-1][0]
    text = "\n".join(
        [
            f"**Sheet:** {sheet.name}",
            "",
            _table(sheet.header, [cells for _, cells in rows]),
        ]
    )
    return Chunk(
        ordinal=ordinal,
        text=text,
        word_count=count_words(text),
        kind=ROW_WINDOW,
        location=f"{sheet.name}, {span_location(ROW_UNIT, start, end)}",
        span_start=start,
        span_end=end,
    )


def _summary_chunk(ordinal: int, title: str, sheet: Sheet) -> Chunk:
    """Assemble the chunk that describes one sheet.

    Args:
        ordinal: The chunk's position in the document.
        title: The workbook's title.
        sheet: The sheet to describe.

    Returns:
        The chunk. It spans nothing, like any chunk written about its
        document rather than cut out of it.
    """
    text = describe_sheet(title, sheet)
    return Chunk(
        ordinal=ordinal,
        text=text,
        word_count=count_words(text),
        kind=SHEET_SUMMARY,
        location=sheet.name,
        span_start=None,
        span_end=None,
    )


def describe_sheet(title: str, sheet: Sheet) -> str:
    """Describe what a sheet holds, from its cells alone.

    Args:
        title: The workbook's title.
        sheet: The sheet.

    Returns:
        A heading line, one sentence about the sheet's size, and one line
        per column saying what is in it.
    """
    count = len(sheet.rows)
    if count:
        first, last = sheet.rows[0][0], sheet.rows[-1][0]
        extent = (
            f"{count} data {'row' if count == 1 else 'rows'} "
            f"({span_location(ROW_UNIT, first, last)})"
        )
    else:
        extent = "no data rows"
    lines = [
        f"**Sheet:** {sheet.name}",
        "",
        f"Sheet {sheet.name!r} of the workbook {title!r} holds {extent} "
        f"under {len(sheet.header)} columns.",
    ]
    if count:
        lines += ["", "Columns:"]
        lines += [
            f"- {label or f'Column {index + 1}'}: "
            f"{_describe_column([row[index] for row in sheet.values], column_format)}"
            for index, (label, column_format) in enumerate(
                zip(sheet.header, sheet.formats, strict=True)
            )
        ]
    return "\n".join(lines)


def _describe_column(values: Sequence[object], number_format: str = "General") -> str:
    """Say what one column holds.

    Numbers and dates are described by their range, text by its values when
    there are few enough of them to be categories and by how many there are
    otherwise. Totals are left out on purpose: a sheet with a total row
    would have that row counted twice, and a column of rates or
    probabilities has no meaningful sum.

    Args:
        values: The column's raw values, one per data row.
        number_format: The column's number format, so a percentage column
            is described in percentages rather than in the fractions they
            are stored as.

    Returns:
        A phrase describing the column.
    """
    present = [value for value in values if value is not None]
    if not present:
        return "empty"
    blanks = len(values) - len(present)
    suffix = f"; {blanks} blank" if blanks else ""

    if all(isinstance(value, datetime.date) for value in present):
        low, high = min(present), max(present)
        return f"dates from {render_cell(low)} to {render_cell(high)}{suffix}"
    if all(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for value in present
    ):
        low = render_cell(min(present), number_format)
        high = render_cell(max(present), number_format)
        return f"numbers from {low} to {high}{suffix}"

    texts = [render_cell(value, number_format) for value in present]
    counts = Counter(texts)
    if len(counts) <= MAX_LISTED_VALUES and len(counts) < len(texts):
        # Most common first; ties in the order they first appear, which
        # Counter preserves and a stable sort keeps.
        ordered = sorted(counts.items(), key=lambda item: -item[1])
        listed = ", ".join(f"{value} ({number})" for value, number in ordered)
        return f"one of {listed}{suffix}"
    return f"text, {len(counts)} distinct values{suffix}"
