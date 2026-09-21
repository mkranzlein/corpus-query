"""Reading a PowerPoint deck and cutting it into one chunk per slide.

This is one format reader among several. It reads a ``.pptx`` file with
``python-pptx`` and cuts it along the seam a deck already has: the slide.

A chunk is one whole slide — its title, the text on it, and its speaker notes
together. The slide itself is usually a fragment ("Margin down 4 points"),
and the sentence that states the claim is the one the presenter says out
loud, which is in the notes. Keeping them in one chunk is what makes the
chunk answerable, and what lets a citation name the slide a reader can turn
to. Slides are not windowed or merged, so ``target_words`` is accepted and
ignored.

Only text is read: titles, text frames (including those inside grouped
shapes and table cells), and notes. Images, charts, and diagrams are not
read, and nothing here attempts a vision pass over them. A slide with no
text at all — no title, no body, no notes — is skipped rather than stored as
an empty chunk. Slide numbers are the deck's own, counted from 1, so a
skipped slide leaves a gap in the numbering rather than shifting every slide
after it.

A deck's title is often not in a title placeholder. Decks built by a script
rather than from a layout commonly draw every piece of text as a free text
box. So a slide's title is taken from its title placeholder when it has one,
and otherwise from the topmost text box set in a heading-sized font that
contains a word. That skips the small kicker line above a cover title and a
bare section number ("01") beside a divider's title.

A deck has one author, read from its core properties, and that author must be
exactly one person on the staff roster. A deck whose author is blank or names
nobody on the roster is refused, because an author the corpus cannot resolve
to a person is worse than none: a question about who wrote something would be
answered with a name nothing else refers to.

Nothing here is random or time-dependent, so reading the same deck twice
gives the same chunks.
"""

from __future__ import annotations

import datetime
import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.util import Pt

from corpus_query.ingest.author import ROSTER_PATH, resolve_author
from corpus_query.ingest.chunk import TARGET_WORDS, Chunk, count_words
from corpus_query.ingest.reader import IngestError, ReadDocument
from corpus_query.store.kinds import PPTX, SLIDE

#: The file extension this reader claims.
DECK_SUFFIX = ".pptx"

#: The smallest font a free text box can be set in and still be taken for a
#: slide's title. Body text and captions sit well under it; slide titles sit
#: well over it.
TITLE_MIN_SIZE = Pt(24)

#: What introduces a slide's speaker notes in its chunk text.
NOTES_LABEL = "Speaker notes:"

_PLACEHOLDER_TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_HAS_LETTER = re.compile(r"[^\W\d_]")


@dataclass(frozen=True)
class ReadSlide:
    """The text of one slide, split into the parts a chunk is built from."""

    number: int
    """The slide's position in the deck, counted from 1."""

    title: str
    """The slide's title, or empty when it has none that can be found."""

    body: tuple[str, ...]
    """Every other line of text on the slide, in shape order."""

    notes: tuple[str, ...]
    """The speaker notes, one entry per non-empty line."""

    @property
    def is_empty(self) -> bool:
        """Whether the slide carries no text at all."""
        return not (self.title or self.body or self.notes)


def read_deck(
    path: Path,
    target_words: int = TARGET_WORDS,
    roster_path: Path | str = ROSTER_PATH,
) -> ReadDocument:
    """Read one PowerPoint deck into a document ready to be written.

    Args:
        path: The deck to read.
        target_words: Accepted to match every other reader, and ignored: a
            deck is chunked by slide, not by size.
        roster_path: The staff roster the deck's author must be on.

    Returns:
        The document, with its author and one chunk per slide that has text.

    Raises:
        IngestError: If the file cannot be opened as a deck, has no date or
            no resolvable author in its core properties, or has no text on
            any slide.
    """
    del target_words
    path = Path(path)
    try:
        presentation = Presentation(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise IngestError(f"Could not open {path} as a PowerPoint deck: {exc}") from exc

    properties = presentation.core_properties
    author = resolve_author(path, properties.author, roster_path)
    document_date = _document_date(properties.modified, properties.created, path)

    slides = [
        read_slide(number, slide)
        for number, slide in enumerate(presentation.slides, start=1)
    ]
    chunks = tuple(chunk_slides(slides))
    if not chunks:
        raise IngestError(f"{path} has no text on any slide.")

    return ReadDocument(
        source_kind=PPTX,
        title=_deck_title(properties.title, slides, path),
        document_date=document_date,
        author=author,
        attendees=(),
        chunks=chunks,
        units=len(slides),
        unit_name="slides",
    )


def read_slide(number: int, slide: Slide) -> ReadSlide:
    """Pull the title, body, and notes text out of one slide.

    Args:
        number: The slide's position in the deck, counted from 1.
        slide: The slide.

    Returns:
        The slide's text. Its title is left out of its body, so it is not
        repeated in the chunk.
    """
    shapes = list(_text_shapes(slide.shapes))
    title_shape = _title_shape(shapes)
    title = " ".join(_lines(title_shape.text_frame.text)) if title_shape else ""
    body = [
        line
        for shape in shapes
        if shape is not title_shape
        for line in _shape_lines(shape)
    ]
    notes: list[str] = []
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            notes = _lines(frame.text)
    return ReadSlide(number=number, title=title, body=tuple(body), notes=tuple(notes))


def chunk_slides(slides: Iterable[ReadSlide]) -> list[Chunk]:
    """Make one chunk per slide that carries any text.

    Args:
        slides: The deck's slides, in order.

    Returns:
        The chunks, in deck order, numbered from 0 with no gaps. Each one's
        span is its slide's number, start and end alike.
    """
    kept = [slide for slide in slides if not slide.is_empty]
    return [_build_chunk(ordinal, slide) for ordinal, slide in enumerate(kept)]


def slide_location(number: int, title: str) -> str:
    """Render a slide's citation label.

    Args:
        number: The slide's number, counted from 1.
        title: The slide's title, possibly empty.

    Returns:
        ``"slide 4 — Ticket History"``, or ``"slide 4"`` for a slide with no
        title.
    """
    if title:
        return f"slide {number} — {title}"
    return f"slide {number}"


def slide_text(slide: ReadSlide) -> str:
    """Render one slide as the text its chunk carries.

    The title comes first as a heading, then the slide's own text, then its
    speaker notes under a label, so that a model or a reader can tell what
    was on the slide from what was said over it. The heading names the slide
    number as well, which keeps slides distinguishable when a whole deck is
    put back together from its chunks.

    Args:
        slide: The slide.

    Returns:
        The chunk text.
    """
    heading = f"## Slide {slide.number}"
    if slide.title:
        heading = f"{heading}: {slide.title}"
    parts = [heading]
    if slide.body:
        parts.append("\n".join(slide.body))
    if slide.notes:
        parts.append("\n".join([NOTES_LABEL, *slide.notes]))
    return "\n\n".join(parts)


def _build_chunk(ordinal: int, slide: ReadSlide) -> Chunk:
    """Assemble one slide's chunk.

    Args:
        ordinal: The chunk's position in the document.
        slide: The slide.

    Returns:
        The chunk.
    """
    text = slide_text(slide)
    return Chunk(
        ordinal=ordinal,
        text=text,
        word_count=count_words(text),
        kind=SLIDE,
        location=slide_location(slide.number, slide.title),
        span_start=slide.number,
        span_end=slide.number,
    )


def _text_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    """Yield every shape that carries text, descending into groups.

    Args:
        shapes: A slide's shapes, or a group's.

    Yields:
        Shapes with a text frame or a table, in the order they are drawn.
    """
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _text_shapes(shape.shapes)
        elif getattr(shape, "has_table", False) or (
            shape.has_text_frame and shape.text_frame.text.strip()
        ):
            yield shape


def _title_shape(shapes: list[BaseShape]) -> BaseShape | None:
    """Find the shape holding a slide's title.

    A title placeholder is the title when there is one. Otherwise the title
    is the topmost text box with a heading-sized font and at least one
    letter in it, with the larger font winning a tie.

    Args:
        shapes: The slide's text-bearing shapes.

    Returns:
        The title shape, or ``None`` when nothing looks like one.
    """
    framed = [shape for shape in shapes if shape.has_text_frame]
    for shape in framed:
        if (
            shape.is_placeholder
            and shape.placeholder_format.type in _PLACEHOLDER_TITLES
        ):
            return shape
    candidates = [
        shape
        for shape in framed
        if _largest_font(shape) >= TITLE_MIN_SIZE
        and _HAS_LETTER.search(shape.text_frame.text)
    ]
    if not candidates:
        return None
    return min(candidates, key=lambda shape: (shape.top or 0, -_largest_font(shape)))


def _largest_font(shape: BaseShape) -> int:
    """Return the largest explicit font size in a shape's text, in EMU.

    Args:
        shape: A shape with a text frame.

    Returns:
        The size, or 0 when no run sets one.
    """
    sizes = [
        run.font.size
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    return max(sizes, default=0)


def _shape_lines(shape: BaseShape) -> list[str]:
    """Return a shape's text as non-empty lines.

    A table is read one row per line, its cells joined with pipes.

    Args:
        shape: A shape with a text frame or a table.

    Returns:
        Its lines, stripped.
    """
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            cells = [" ".join(_lines(cell.text)) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return rows
    return _lines(shape.text_frame.text)


def _lines(text: str) -> list[str]:
    """Split text into stripped, non-empty lines.

    ``python-pptx`` renders a soft line break inside a paragraph as a
    vertical tab, which is a line break for our purposes too.

    Args:
        text: The text to split.

    Returns:
        Its lines.
    """
    return [
        line.strip() for line in text.replace("\v", "\n").splitlines() if line.strip()
    ]


def _document_date(
    modified: datetime.datetime | None,
    created: datetime.datetime | None,
    path: Path,
) -> str:
    """Pick the date a deck is filed under.

    Args:
        modified: When the core properties say it was last modified.
        created: When they say it was created.
        path: The deck, for the error message.

    Returns:
        The last-modified date, or the created date when there is none, as
        ``YYYY-MM-DD``.

    Raises:
        IngestError: If the deck carries neither.
    """
    when = modified or created
    if when is None:
        raise IngestError(f"{path} has no created or modified date in its properties.")
    return when.date().isoformat()


def _deck_title(title: str | None, slides: list[ReadSlide], path: Path) -> str:
    """Pick a deck's title.

    Args:
        title: The title in the deck's core properties.
        slides: The deck's slides.
        path: The deck, used as the last resort.

    Returns:
        The core-properties title, else the first slide title found, else the
        file's name.
    """
    if title and title.strip():
        return title.strip()
    for slide in slides:
        if slide.title:
            return slide.title
    return path.stem
